#!/usr/bin/env python3
"""Build the hackathon deck for Cognitive Fitness & Human Development.

Usage:  python3 build_slides.py [output.pptx]   (default: same directory)
Requires python-pptx (`pip install python-pptx`). Slide copy lives here so
the deck can be regenerated from the repo; see SLIDES.md for the outline.
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt
from lxml import etree

INK = RGBColor(0x1C, 0x2B, 0x2F)
TEAL = RGBColor(0x0E, 0x6E, 0x66)
MUTED = RGBColor(0x5A, 0x6A, 0x6E)
BG = RGBColor(0xF7, 0xF5, 0xEF)
CARD = RGBColor(0xEC, 0xEF, 0xE9)
ACCENT2 = RGBColor(0xB4, 0x53, 0x09)  # burnt orange for warnings/arrows out
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def box(s, l, t, w, h, fill=None, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def text(s, l, t, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, leading=1.0):
    """runs: str, or list of (str, dict) paragraphs; dict may set size/color/bold/bullet/level."""
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, {})]
    for i, (txt, o) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = o.get("level", 0)
        p.alignment = o.get("align", align)
        if o.get("space_before"):
            p.space_before = Pt(o["space_before"])
        if leading != 1.0:
            p.line_spacing = leading
        if o.get("bullet"):
            txt = "•  " + txt
        r = p.add_run()
        r.text = txt
        f = r.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
        f.name = "Calibri"
    return tb


def header(s, title, kicker=None):
    box(s, 0, 0, W, Inches(0.16), fill=TEAL)
    if kicker:
        text(s, Inches(0.6), Inches(0.32), W - Inches(1.2), Inches(0.35),
             kicker.upper(), size=13, color=MUTED, bold=True)
    text(s, Inches(0.6), Inches(0.62), W - Inches(1.2), Inches(0.85),
         title, size=30, color=INK, bold=True)


def note(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def flowbox(s, l, t, w, h, title, sub=None, fill=CARD, line=TEAL, tcolor=INK, size=15):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.4)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = tcolor; r.font.name = "Calibri"
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(size - 4); r2.font.color.rgb = MUTED; r2.font.name = "Calibri"
    return sh


def arrow(s, x1, y1, x2, y2, color=TEAL, dashed=False, width=1.8):
    from pptx.oxml.ns import qn
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    spPr = conn._element.find(qn("p:spPr"))
    ln = spPr.find(qn("a:ln"))
    if dashed:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    # child order inside a:ln is fill -> dash -> headEnd -> tailEnd; appending after
    # the solidFill/width already written by python-pptx satisfies that.
    head = etree.SubElement(ln, qn("a:headEnd"))
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("len", "med")
    return conn


# 1 — Title ---------------------------------------------------------------
s = slide()
box(s, 0, 0, W, Inches(0.22), fill=TEAL)
text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.6),
     "Cognitive Fitness & Human Development", size=48, bold=True, color=INK)
text(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(0.7),
     "Evidence-informed learning for how people think, learn, focus, sleep, and connect.",
     size=22, color=MUTED)
text(s, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.6),
     "“Every page earns its place” — purpose, structure, linkage, and direction, enforced by the build.",
     size=18, color=TEAL, bold=True)
text(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.9),
     [("Hackathon submission · October 2026", {"size": 14, "color": MUTED}),
      ("github.com/ahmed-raza-ur-rehman/hackathon2  ·  cognitive-fitness-blond.vercel.app",
       {"size": 14, "color": MUTED, "space_before": 4})])
note(s, "[0:00] This is Cognitive Fitness and Human Development — an evidence-informed learning "
        "space for how people think, learn, focus, sleep, and build habits. The problem it attacks: "
        "brain-training products overclaim, learning tools are scattered, and most interfaces bury "
        "the next step under decoration.")

# 2 — Problem -------------------------------------------------------------
s = slide()
header(s, "The problem", "01 · problem statement")
items = [
    ("Overclaiming", "“Brain training” apps promise transfer effects the evidence doesn't support "
     "(Simons et al. 2016 meta-review is one we cite in-app)."),
    ("Fragmented learning", "Concepts, study technique, and practice live in different products with no "
     "guided path from understanding to application."),
    ("Cognitive overload", "Learning UIs duplicate navigation, bury the next action, and ignore reading "
     "preferences (size, contrast, motion)."),
    ("AI demos that exclude", "Education demos hide an operator-billed API key behind logins — or run up "
     "the developer's bill per request."),
]
runs = []
for t, d in items:
    runs.append((t, {"size": 21, "bold": True, "color": TEAL, "space_before": 10}))
    runs.append((d, {"size": 17, "color": INK, "level": 0}))
text(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), runs)
note(s, "[0:25] Research — Simons and colleagues' 2016 review is one example we cite — shows game-based "
        "training rarely transfers. Yet most apps still sell 'your brain, now at 20 percent better'.")

# 3 — Evidence stance -----------------------------------------------------
s = slide()
header(s, "We publish our own limits", "02 · the honesty stance")
runs = [
    ("Every page carries a Boundary", {"size": 20, "bold": True, "color": TEAL}),
    ("Scores stay task-specific — never a diagnosis, never IQ, never proof of broad cognitive change.",
     {"size": 17, "space_before": 2}),
    ("An inspectable evidence ledger", {"size": 20, "bold": True, "color": TEAL, "space_before": 14}),
    ("Claims link to studies with a plain-language “what it does/doesn't mean” summary — "
     "users can check us.", {"size": 17, "space_before": 2}),
    ("Careful copy rules", {"size": 20, "bold": True, "color": TEAL, "space_before": 14}),
    ("Adult-support language for emotion topics; clinical content explicitly excluded; "
     "four-week schedules are authored routines, not validated curricula.", {"size": 17, "space_before": 2}),
]
text(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(5.2), runs)
note(s, "[0:25 cont.] So every page in this app must carry a boundary: what the task score does NOT prove. "
        "You can see it on any concept or evidence page.")

# 4 — The contract --------------------------------------------------------
s = slide()
header(s, "Structure as code, not taste", "03 · problem → solution")
runs = [
    ("Every route is registered in content/routes-v1.json:", {"size": 19, "bold": True}),
    ("kind · audience · purpose · one primary action · anchors · leadsTo", {"size": 18, "color": TEAL}),
    ("A build-time validator (content/validate.mjs) fails the build if any page is unregistered, "
     "dangling, or lacks a purpose.", {"size": 17, "space_before": 8}),
    ("The page shell renders breadcrumbs, the stated purpose, a boundary note, and a registry-driven "
     "“where this leads” panel — from the same data.", {"size": 17, "space_before": 6}),
    ("23 registered pages · 6 guided journeys · 29 linked steps — each step with purpose, action, "
     "takeaway, and caution.", {"size": 17, "color": ACCENT2, "bold": True, "space_before": 10}),
]
text(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(4.3), runs)
y = Inches(6.15); hh = Inches(0.75); ww = Inches(2.6); gap = Inches(0.75)
xs = [Inches(0.9) + i * (ww + gap) for i in range(4)]
labels = [("content/routes-v1.json", "single registry"),
          ("content/validate.mjs", "purpose + links check"),
          ("CI — lint · tsc · 187 tests", "hosted, on every push"),
          ("PageIntro shell", "renders structure on-page")]
for x, (t, sub) in zip(xs, labels):
    flowbox(s, x, y, ww, hh, t, sub, size=13)
for a, b in zip(xs, xs[1:]):
    arrow(s, a + ww, y + hh / 2, b, y + hh / 2)
note(s, "[1:05] Our solution is structure you can verify. Every one of the twenty-three pages is registered "
        "with an audience, a purpose, one primary action, and links onward. The validator fails the build if "
        "any of that is missing — 'every page has meaning' is enforced by the pipeline, not by taste.")

# 5 — What we built -------------------------------------------------------
s = slide()
header(s, "What we built", "04 · solution")
feats = [
    ("Guided journeys", "6 paths, 29 steps: understand → experience → inspect → practice → apply → reflect. "
     "No account, no tracking."),
    ("Client-side search", "Finds concepts, steps, activities, claims, sources; typo-grace; works with zero JS."),
    ("7 reading settings", "Size, contrast, spacing, serif, motion, reading guide, underlines — pre-paint, "
     "cross-tab sync, storage-failure tolerant."),
    ("Local-only progress", "Describes opted-in practice sessions as counts and cadence; no pooled score."),
    ("3-mode coach", "Device-only default · Puter user-pays opt-in · operator path disabled by default."),
    ("Atlas & practice", "14 development areas with starters and boundaries; 8 untimed practice modes; "
     "4 four-week self-owned schedules."),
]
runs = []
for t, d in feats:
    runs.append((t, {"size": 18, "bold": True, "color": TEAL}))
    runs.append((d, {"size": 15.5, "space_before": 1}))
text(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.6), runs)
note(s, "[1:05 cont.] On top of the contract: six guided journeys that chain understand, experience, inspect, "
        "practice, apply, and reflect — with no account and no tracking.")

# 6 — User-pays AI --------------------------------------------------------
s = slide()
header(s, "AI without operator keys: user-pays", "05 · the Puter model")
runs = [
    ("Each visitor pays for their own usage", {"size": 19, "bold": True, "color": TEAL}),
    ("Puter.js runs in the browser; the visitor signs into their own Puter account. Every user starts with a "
     "free monthly allowance — then they pay. We never hold credentials, and our dev cost is $0.",
     {"size": 16}),
    ("Guardrails, shown in the UI", {"size": 19, "bold": True, "color": TEAL, "space_before": 12}),
    ("Consent asked per message, re-asked when context changes · one bounded turn per request · "
     "6 requests / 5 min per page · adult self-declaration (not verification) · kill switch env flag · "
     "fail-closed: provider down ⇒ the page still works.", {"size": 16, "space_before": 2}),
    ("Honest framing", {"size": 19, "bold": True, "color": ACCENT2, "space_before": 12}),
    ("“Free monthly allowance, then paid” — the product says so in the consent box. No unlimited-free-AI "
     "claims, and Google sign-in is deliberately NOT wired as a model switch (auth ≠ model authorization).",
     {"size": 16, "space_before": 2}),
]
text(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), runs)
note(s, "[4:10 on camera] Opt in, and it runs on user-pays billing: each visitor gets a free monthly "
        "allowance, then they pay their own usage. We never hold credentials; every message needs explicit "
        "consent; one kill switch turns it off.")

# 7 — Tech stack ----------------------------------------------------------
s = slide()
header(s, "Technology stack", "06 · tools & technologies")
runs = [
    ("App", {"size": 18, "bold": True, "color": TEAL}),
    ("Next.js 16.3.8 (App Router, static-first) · React 19.3 · TypeScript 5.9 strict · "
     "Tailwind 4 + CSS Modules", {"size": 16}),
    ("Content as data", {"size": 18, "bold": True, "color": TEAL, "space_before": 10}),
    ("JSON route registry + journey content, custom Node ≥22 validators wired into CI", {"size": 16}),
    ("Quality", {"size": 18, "bold": True, "color": TEAL, "space_before": 10}),
    ("Vitest 5 + Testing Library + jsdom (187 tests) · ESLint 9 · hosted Actions CI · Vercel builds", {"size": 16}),
    ("Hosting & AI", {"size": 18, "bold": True, "color": TEAL, "space_before": 10}),
    ("Vercel (Node 22, root cognitive-fitness) · Puter.js SDK client-side (default model gpt-5.4-nano) · "
     "OpenAI server route present but fail-closed without a key", {"size": 16}),
    ("Security posture", {"size": 18, "bold": True, "color": TEAL, "space_before": 10}),
    ("CSP/SRI for the third-party script is a tracked open decision (R-04), claimed nowhere as done", {"size": 16}),
]
text(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), runs)
note(s, "[1:45] The stack: Next.js 16 with React 19, strict TypeScript, content as JSON data with custom "
        "validators, Vitest with 187 tests, Tailwind and CSS modules, deployed on Vercel, and Puter.js for "
        "the optional AI guide.")

# 8 — High-level diagram --------------------------------------------------
s = slide()
header(s, "High-level architecture", "07 · diagram")
yy0, hh0 = Inches(2.0), Inches(0.95)
browser = flowbox(s, Inches(0.7), yy0, Inches(2.9), hh0, "Visitor browser", "one localStorage key · no cookies")
cdn = flowbox(s, Inches(4.7), yy0, Inches(3.2), hh0, "Vercel CDN", "22 prerendered pages (static)")
ls = flowbox(s, Inches(0.7), yy0 + Inches(1.55), Inches(2.9), Inches(0.8),
             "localStorage", "preferences · saved practice")
optin = flowbox(s, Inches(4.7), yy0 + Inches(1.7), Inches(3.2), Inches(0.9),
                "Coach panel — opt-in", "mode chooser · consent · rate limit")
puter = flowbox(s, Inches(9.0), yy0 + Inches(0.15), Inches(3.6), Inches(0.9),
                "Puter.js (in browser)", "user signs in · free allowance → user pays")
model = flowbox(s, Inches(9.0), yy0 + Inches(1.55), Inches(3.6), Inches(0.9),
                "Provider model", "e.g. gpt-5.4-nano · app holds NO key")
api = flowbox(s, Inches(9.0), yy0 + Inches(2.95), Inches(3.6), Inches(0.85),
              "/coach/api (OpenAI)", "disabled: no key configured → 503",
              fill=BG, line=MUTED, tcolor=MUTED)
pr = flowbox(s, Inches(0.7), Inches(5.9), Inches(2.9), Inches(0.85), "GitHub PR", "arena branch → main")
ci = flowbox(s, Inches(4.7), Inches(5.9), Inches(3.2), Inches(0.85),
             "CI gate", "lint · tsc · 187 tests · validators · build")
dep = flowbox(s, Inches(9.0), Inches(5.9), Inches(3.6), Inches(0.85),
              "Vercel deploy", "preview on PR · prod on main")
arrow(s, Inches(3.6), yy0 + hh0 / 2, Inches(4.7), yy0 + hh0 / 2)
arrow(s, Inches(2.15), yy0 + Inches(0.95), Inches(2.15), yy0 + Inches(1.55))
arrow(s, Inches(6.3), yy0 + Inches(1.7), Inches(2.6), yy0 + Inches(1.55) + Inches(0.8), dashed=True, color=ACCENT2)
arrow(s, Inches(7.9), yy0 + Inches(0.62), Inches(9.0), yy0 + Inches(0.62))
arrow(s, Inches(10.8), yy0 + Inches(1.05), Inches(10.8), yy0 + Inches(1.55))
arrow(s, Inches(7.9), yy0 + Inches(2.15), Inches(9.0), yy0 + Inches(3.35), dashed=True, color=MUTED)
arrow(s, Inches(3.6), Inches(6.32), Inches(4.7), Inches(6.32))
arrow(s, Inches(7.9), Inches(6.32), Inches(9.0), Inches(6.32))
arrow(s, Inches(10.8), Inches(5.9), Inches(10.8), yy0 + Inches(2.2), dashed=True, color=MUTED)
text(s, Inches(4.7), Inches(6.9), Inches(7.9), Inches(0.5),
     "Solid = always-on paths · orange dashed = user data never leaves the device unless AI is opted into · "
     "grey = disabled unless configured", size=12, color=MUTED)
note(s, "[2:20] The browser talks to a static CDN — 22 prerendered pages. Preferences and history live only "
        "in local storage. Opt into the guide and your browser calls the provider through Puter — no key in "
        "our infra. CI gates every deploy.")

# 9 — Privacy -------------------------------------------------------------
s = slide()
header(s, "Privacy is a design constraint, not a policy page", "08 · trust")
pruns = [
    ("Zero accounts, zero databases, zero analytics, zero ads.", {}),
    ("One localStorage key holds reading preferences and (opt-in) saved practice history.", {}),
    ("Search runs entirely in the browser — no queries, no logs.", {}),
    ("Challenge/result codes are self-reported by design; the UI says so — no fake verification claims.", {}),
    ("/privacy lists every data path in plain language — including what Puter's own terms cover once you sign in.", {}),
    ("Security headers shipped; CSP/SRI decisions for the third-party script are tracked, not hand-waved.",
     {"color": ACCENT2}),
]
text(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(5.0),
     [(t, {**o, "bullet": True, "space_before": 8}) for t, o in pruns])
note(s, "[4:35] No accounts, no database, no analytics — the privacy page lists every path data takes, "
        "because this is a research prototype with honest gates, not a launch claim.")

# 10 — Verified & honest --------------------------------------------------
s = slide()
header(s, "What is verified — and what isn't", "09 · quality evidence")
box(s, Inches(0.6), Inches(1.7), Inches(5.95), Inches(5.0), fill=CARD)
box(s, Inches(6.8), Inches(1.7), Inches(5.95), Inches(5.0), fill=BG, line=ACCENT2)
text(s, Inches(0.95), Inches(1.95), Inches(5.3), Inches(4.6),
     [("Verified today", {"size": 19, "bold": True, "color": TEAL, "space_before": 0}),
      ("187 unit/behavior tests green (37 files)", {"size": 15.5, "space_before": 10, "bullet": True}),
      ("ESLint 0 errors · tsc strict clean", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("Registry + journey validators pass (23 pages, 6/29)", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("Hosted CI green on push + PR (Actions) and a Vercel build", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("22/23 routes prerendered static; headers verified", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("Orchestrated audit trail (validate + 11 workflow tests)", {"size": 15.5, "space_before": 6, "bullet": True})])
text(s, Inches(7.15), Inches(1.95), Inches(5.3), Inches(4.6),
     [("Still open — deliberately unchecked", {"size": 19, "bold": True, "color": ACCENT2}),
      ("Manual a11y/device matrix (screen reader, keyboard, zoom, touch)",
       {"size": 15.5, "space_before": 10, "bullet": True}),
      ("Legal/privacy review — /privacy is a prototype disclosure", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("Independent security + evidence review", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("CSP/SRI decision for Puter script (R-04)", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("Adversarial prompting of the AI guide (R-05)", {"size": 15.5, "space_before": 6, "bullet": True}),
      ("main branch protection: require the checks (owner action)", {"size": 15.5, "space_before": 6, "bullet": True})])
note(s, "[4:35 cont.] Nothing here is claimed WCAG-conformant or production-ready — that restraint is the "
        "point: the checklist in the repo lists exactly which gates remain.")

# 11 — Finance ------------------------------------------------------------
s = slide()
header(s, "Finance: $0/month by construction", "10 · economics")
rows = [
    ("Development toolchain", "$0", "open-source stack; local dev + CI free tier"),
    ("Hosting", "$0", "Vercel Hobby (Hobby forbids commercial use → Pro $20/dev/mo if it productizes)"),
    ("Domain", "$0", "vercel.app subdomain; optional custom ≈ $10–20/yr"),
    ("AI — user-pays mode", "$0 to us", "each user: free monthly allowance → then their own metered usage; no key custody"),
    ("AI — operator path", "$0 (disabled)", "if enabled, operator pays per token; request hard-bounded (~2k in / ≤900 chars out)"),
    ("Analytics / email / ads", "$0", "deliberately none — a privacy decision, not a budget one"),
]
yy = Inches(1.8)
text(s, Inches(0.8), yy - Inches(0.35), Inches(11), Inches(0.3),
     "Item · Cost today · Note", size=13, color=MUTED, bold=True)
for item, cost, note_ in rows:
    text(s, Inches(0.8), yy, Inches(3.6), Inches(0.6), item, size=16, bold=True)
    text(s, Inches(4.5), yy, Inches(1.9), Inches(0.6), cost, size=16, color=TEAL, bold=True)
    text(s, Inches(6.5), yy, Inches(6.2), Inches(0.7), note_, size=14, color=MUTED)
    yy += Inches(0.72)
text(s, Inches(0.8), Inches(6.35), Inches(11.8), Inches(0.7),
     "Risk the checklist tracks: a public “AI product” on Hobby is the one combo that forces a plan decision "
     "— flagged for the owner, priced, not promised.", size=14, color=ACCENT2)
note(s, "Finance detail is PRD § 9: total operating cost today $0; the only scaling costs are hosting tier "
        "and operator-run OpenAI, both gated on explicit owner decisions.")

# 12 — Roadmap & asks -----------------------------------------------------
s = slide()
header(s, "Roadmap & what we need next", "11 · after the hackathon")
runs = [
    ("Ship-safe (blocking, in the repo checklist)", {"size": 18, "bold": True, "color": TEAL}),
    ("Run the accessibility/device matrix · legal review · CSP/SRI ruling · adversarial AI tests · "
     "protect main with the required check", {"size": 16}),
    ("Product growth (approved direction)", {"size": 18, "bold": True, "color": TEAL, "space_before": 12}),
    ("More journeys from the same registry · teacher-mode handouts · concept-to-schedule deep links · "
     "multi-language content pipeline", {"size": 16}),
    ("Explicitly not on the near roadmap", {"size": 18, "bold": True, "color": ACCENT2, "space_before": 12}),
    ("Accounts, streaks-as-pressure, public rankings, diagnosis-adjacent features, “brain score”",
     {"size": 16}),
]
text(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(5.0), runs)
note(s, "Asks: 30 minutes of judge time for the manual matrix on one device; a Vercel plan confirmation; "
        "and one education reviewer to pressure-test the boundary copy.")

# 13 — Demo & close -------------------------------------------------------
s = slide()
box(s, 0, 0, W, Inches(0.22), fill=TEAL)
text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.0),
     "Live demo", size=44, bold=True, color=INK)
runs = [
    ("Start Here journey → one step, its caution, its link", {"size": 19, "bullet": True}),
    ("Search “spac” → spacing effect, zero requests", {"size": 19, "bullet": True, "space_before": 8}),
    ("Reading settings → serif + reading guide, instant", {"size": 19, "bullet": True, "space_before": 8}),
    ("/progress → opt-in insight, no pooled score", {"size": 19, "bullet": True, "space_before": 8}),
    ("/coach → three modes; user-pays guide", {"size": 19, "bullet": True, "space_before": 8}),
]
text(s, Inches(1.1), Inches(3.1), Inches(11.2), Inches(3.0), runs)
text(s, Inches(1.1), Inches(6.2), Inches(11.2), Inches(0.9),
     [("Thank you — every word on every page is link-checked by CI.", {"size": 18, "color": TEAL, "bold": True}),
      ("github.com/ahmed-raza-ur-rehman/hackathon2 · cognitive-fitness-blond.vercel.app",
       {"size": 14, "color": MUTED, "space_before": 6})])
note(s, "[2:50 → 5:00] Demo per the scene table; closing line: 'No accounts, no database, no analytics — "
        "a research prototype with honest gates, not a launch claim. Thank you.'")

out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "cognitive-fitness-slides.pptx"
prs.save(str(out))
print(f"wrote {out} ({len(prs.slides._sldIdLst)} slides)")
